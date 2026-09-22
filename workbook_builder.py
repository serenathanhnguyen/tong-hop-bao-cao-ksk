"""Dựng file Excel "Bảng tổng hợp phân loại sức khỏe" theo đúng khung mẫu
mô tả trong docs/QUY_TAC.md (mục 1-6)."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import config
from .classify import PersonDerived
from .raw_reader import Person, parse_vn_number

THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")
BOLD = Font(bold=True)
BOLD_CENTER_FONT = Font(bold=True)


@dataclass
class ReportMeta:
    ten_phuong: str = "PHƯỜNG NHIÊU LỘC"
    ten_tram_yte: str = "TRẠM Y TẾ"
    doi_tuong: str = ""  # vd "CB-NV CÔNG TY TNHH CÁ SẤU VIỆT PHONG"
    nam: str = ""
    nguoi_lap_bang: str = config.DEFAULT_NGUOI_LAP_BANG
    giam_doc: str = ""  # để trống mặc định, ký tay
    ngay_ky: date | None = None  # None -> để trống dạng "Ngày ... tháng ... năm ..."


@dataclass
class ColumnDef:
    key: str
    label: str
    group: str | None = None  # "mat" | "lam_sang" | "can_lam_sang" | None


def build_column_defs(active_cls_cols: list[dict]) -> list[ColumnDef]:
    cols: list[ColumnDef] = [
        ColumnDef("stt", "STT"),
        ColumnDef("ho_ten", "HỌ & TÊN"),
        ColumnDef("ten", "TÊN"),
        ColumnDef("cccd", "CCCD"),
        ColumnDef("ngay_sinh", "NGÀY SINH"),
        ColumnDef("gioi_tinh", "GIỚI TÍNH"),
        ColumnDef("chieu_cao", "Chiều cao (Cm)"),
        ColumnDef("can_nang", "Cân nặng (Kg)"),
        ColumnDef("bmi", "BMI"),
        ColumnDef("huyet_ap", "Huyết áp\n(mmHg)"),
        ColumnDef("mat_p", "P", group="mat"),
        ColumnDef("mat_t", "T", group="mat"),
    ]
    for label in config.TEMPLATE_CLINICAL_COLS:
        cols.append(ColumnDef(f"ls_{label}", label, group="lam_sang"))
    for c in active_cls_cols:
        cols.append(ColumnDef(f"cls_{c['key']}", c["label"], group="can_lam_sang"))
    cols.append(ColumnDef("xep_loai", "Xếp loại"))
    cols.append(ColumnDef("ghi_chu", "GHI CHÚ"))
    cols.append(ColumnDef("canh_bao", "CẢNH BÁO\n(để kiểm tra lại, xóa cột\nsau khi đã rà soát)"))
    return cols


ROW_TITLE1, ROW_TITLE2 = 1, 2
ROW_TITLE_BLANK = 3
ROW_TITLE4, ROW_TITLE5 = 4, 5
ROW_HEADER_GROUP, ROW_HEADER_SUB = 7, 8
DATA_START_ROW = 9


def _fmt_ngay_sinh(v: Any) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    return str(v).strip()


def _set(ws: Worksheet, row: int, col: int, value: Any, *, font: Font | None = None,
         align: Alignment | None = None, border: Border | None = BOX):
    cell = ws.cell(row=row, column=col, value=value)
    # QUY TẮC QUAN TRỌNG 3: luôn set lại font tường minh, không dựa vào mặc định của ô.
    cell.font = font if font is not None else Font(bold=False, italic=False)
    if align is not None:
        cell.alignment = align
    if border is not None:
        cell.border = border
    return cell


def build_workbook(
    meta: ReportMeta,
    people: list[Person],
    derived: list[PersonDerived],
    active_cls_cols: list[dict],
    ghi_chu_list: list[str],
    canh_bao_list: list[str],
) -> Workbook:
    cols = build_column_defs(active_cls_cols)
    ncols = len(cols)
    last_col_idx = ncols
    last_col_letter = get_column_letter(last_col_idx)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # ---- Mục 1: tiêu đề ----------------------------------------------------
    ws.merge_cells(f"A{ROW_TITLE1}:G{ROW_TITLE1}")
    _set(ws, ROW_TITLE1, 1, "ỦY BAN NHÂN DÂN " + meta.ten_phuong.upper(), font=BOLD, align=CENTER, border=None)
    ws.merge_cells(f"T{ROW_TITLE1}:{last_col_letter}{ROW_TITLE1}")
    _set(ws, ROW_TITLE1, 20, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", font=BOLD, align=CENTER, border=None)

    ws.merge_cells(f"A{ROW_TITLE2}:G{ROW_TITLE2}")
    _set(ws, ROW_TITLE2, 1, meta.ten_tram_yte.upper(), font=BOLD, align=CENTER, border=None)
    ws.merge_cells(f"T{ROW_TITLE2}:{last_col_letter}{ROW_TITLE2}")
    _set(ws, ROW_TITLE2, 20, "Độc lập - Tự do - Hạnh phúc", font=Font(bold=True, italic=False), align=CENTER, border=None)

    doi_tuong = meta.doi_tuong.strip() or "..."
    ws.merge_cells(f"A{ROW_TITLE4}:{last_col_letter}{ROW_TITLE4}")
    _set(ws, ROW_TITLE4, 1, f"BẢNG TỔNG HỢP PHÂN LOẠI SỨC KHỎE {doi_tuong}".upper(),
         font=Font(bold=True, size=13), align=CENTER, border=None)

    nam = meta.nam.strip() or "..."
    ws.merge_cells(f"A{ROW_TITLE5}:{last_col_letter}{ROW_TITLE5}")
    _set(ws, ROW_TITLE5, 1, f"KHÁM SỨC KHỎE ĐỊNH KỲ NĂM {nam}".upper(),
         font=Font(bold=True, size=12), align=CENTER, border=None)

    # ---- Mục 2: header cột (row 7-8) ---------------------------------------
    group_spans: dict[str, list[int]] = {}
    for idx, c in enumerate(cols, start=1):
        if c.group:
            group_spans.setdefault(c.group, []).append(idx)

    GROUP_LABELS = {"mat": "MẮT 10/10", "lam_sang": "KHÁM LÂM SÀNG", "can_lam_sang": "CẬN LÂM SÀNG"}

    for idx, c in enumerate(cols, start=1):
        letter = get_column_letter(idx)
        if c.group:
            span = group_spans[c.group]
            if idx == span[0]:
                first_l, last_l = get_column_letter(span[0]), get_column_letter(span[-1])
                ws.merge_cells(f"{first_l}{ROW_HEADER_GROUP}:{last_l}{ROW_HEADER_GROUP}")
                _set(ws, ROW_HEADER_GROUP, span[0], GROUP_LABELS[c.group], font=BOLD, align=CENTER)
                for j in span[1:]:
                    _set(ws, ROW_HEADER_GROUP, j, None, align=CENTER)
            _set(ws, ROW_HEADER_SUB, idx, c.label, font=BOLD, align=CENTER)
        else:
            ws.merge_cells(f"{letter}{ROW_HEADER_GROUP}:{letter}{ROW_HEADER_SUB}")
            _set(ws, ROW_HEADER_GROUP, idx, c.label, font=BOLD, align=CENTER)
            _set(ws, ROW_HEADER_SUB, idx, None, align=CENTER)

    # ---- Mục 3: vùng dữ liệu -----------------------------------------------
    col_index = {c.key: i + 1 for i, c in enumerate(cols)}

    def col(key: str) -> int:
        return col_index[key]

    for i, (p, d) in enumerate(zip(people, derived)):
        row = DATA_START_ROW + i
        ghi_chu, canh_bao = ghi_chu_list[i], canh_bao_list[i]

        _set(ws, row, col("stt"), i + 1, align=CENTER)
        _set(ws, row, col("ho_ten"), p.ho_ten, align=Alignment(horizontal="left", vertical="center"))
        ten_rieng = p.ho_ten.strip().split(" ")[-1] if p.ho_ten.strip() else ""
        _set(ws, row, col("ten"), ten_rieng, align=CENTER)
        _set(ws, row, col("cccd"), p.cccd, align=CENTER)
        _set(ws, row, col("ngay_sinh"), _fmt_ngay_sinh(p.ngay_sinh), align=CENTER)
        _set(ws, row, col("gioi_tinh"), p.gioi_tinh, align=CENTER)
        _set(ws, row, col("chieu_cao"), p.chieu_cao, align=CENTER)
        _set(ws, row, col("can_nang"), p.can_nang, align=CENTER)
        g_letter, h_letter = get_column_letter(col("chieu_cao")), get_column_letter(col("can_nang"))
        _set(ws, row, col("bmi"), f"={h_letter}{row}/({g_letter}{row}*{g_letter}{row})*10000", align=CENTER)
        _set(ws, row, col("huyet_ap"), p.huyet_ap_text, align=CENTER)
        _set(ws, row, col("mat_p"), p.mat_phai, align=CENTER)
        _set(ws, row, col("mat_t"), p.mat_trai, align=CENTER)

        for label in config.TEMPLATE_CLINICAL_COLS:
            val = d.clinical_cols.get(label)
            _set(ws, row, col(f"ls_{label}"), val if val is not None else "", align=CENTER)

        for c in active_cls_cols:
            key = f"cls_{c['key']}"
            kind = c["kind"]
            if kind == "cbc":
                _set(ws, row, col(key), d.cbc_result, align=CENTER)
            elif kind == "chem":
                flag = d.lab_flags[c["label"]]
                font = Font(bold=flag.bold, italic=flag.italic)
                _set(ws, row, col(key), flag.value, font=font, align=CENTER)
            elif kind == "urinalysis":
                _set(ws, row, col(key), d.urinalysis_result, align=CENTER)
            elif kind == "us_general":
                _set(ws, row, col(key), d.us_general_result, align=CENTER)
            elif kind == "us_breast":
                _set(ws, row, col(key), d.us_breast_result, align=CENTER)

        _set(ws, row, col("xep_loai"), d.xep_loai if d.xep_loai is not None else "", font=BOLD, align=CENTER)
        _set(ws, row, col("ghi_chu"), ghi_chu, align=LEFT_WRAP)
        _set(ws, row, col("canh_bao"), canh_bao, font=Font(italic=True), align=LEFT_WRAP)

    last_data_row = DATA_START_ROW + len(people) - 1 if people else DATA_START_ROW

    # ---- Mục 4: dòng TỔNG CỘNG ----------------------------------------------
    tong_cong_row = last_data_row + 1
    ws.merge_cells(f"A{tong_cong_row}:E{tong_cong_row}")
    _set(ws, tong_cong_row, 1, "TỔNG CỘNG", font=BOLD, align=CENTER)
    if people:
        for c in cols:
            if c.key in ("stt", "ho_ten", "ten", "cccd", "ngay_sinh", "ghi_chu", "canh_bao"):
                continue
            idx = col_index[c.key]
            letter = get_column_letter(idx)
            _set(ws, tong_cong_row, idx,
                 f"=COUNTA({letter}{DATA_START_ROW}:{letter}{last_data_row})",
                 font=BOLD, align=CENTER)

    # ---- Mục 5: khối thống kê xếp loại --------------------------------------
    stats_start = tong_cong_row + 3
    xep_loai_letter = get_column_letter(col_index["xep_loai"])
    _set(ws, stats_start, 1, "Tổng số:", font=BOLD, align=Alignment(horizontal="left"), border=None)
    _set(ws, stats_start, 3, f"=SUM(C{stats_start + 1}:C{stats_start + 5})", align=CENTER, border=None)
    _set(ws, stats_start, 4, "Người", align=Alignment(horizontal="left"), border=None)
    for i in range(1, 6):
        r = stats_start + i
        _set(ws, r, 1, f"Loại {i}:", align=Alignment(horizontal="left"), border=None)
        if people:
            _set(ws, r, 3,
                 f'=COUNTIF(${xep_loai_letter}${DATA_START_ROW}:${xep_loai_letter}${last_data_row},"{i}")',
                 align=CENTER, border=None)
        else:
            _set(ws, r, 3, 0, align=CENTER, border=None)
        _set(ws, r, 4, "Người", align=Alignment(horizontal="left"), border=None)

    note_row = stats_start + 6
    ws.merge_cells(f"A{note_row}:{last_col_letter}{note_row}")
    _set(ws, note_row, 1,
         "Ghi chú: chỉ số cận lâm sàng in đậm = cao hơn giá trị tham chiếu; "
         "in nghiêng = thấp hơn giá trị tham chiếu.",
         font=Font(italic=True), align=Alignment(horizontal="left"), border=None)

    # ---- Mục 6: phần ký tên --------------------------------------------------
    sign_col_start = max(1, last_col_idx - 6)
    sign_col_letter_start = get_column_letter(sign_col_start)

    date_row = note_row + 2
    ws.merge_cells(f"{sign_col_letter_start}{date_row}:{last_col_letter}{date_row}")
    if meta.ngay_ky is not None:
        ngay_text = (f"Ngày {meta.ngay_ky.day} tháng {meta.ngay_ky.month} "
                     f"năm {meta.ngay_ky.year}")
    else:
        # QUY TẮC QUAN TRỌNG 4: giữ dạng chấm chấm để trống, không tự điền ngày.
        ngay_text = "Ngày ......... tháng ......... năm ........."
    _set(ws, date_row, sign_col_start, ngay_text, font=Font(italic=True), align=RIGHT, border=None)

    label_row = date_row + 1
    _set(ws, label_row, 2, "Người Lập Bảng", font=BOLD, align=Alignment(horizontal="center"), border=None)
    ws.merge_cells(f"{sign_col_letter_start}{label_row}:{last_col_letter}{label_row}")
    _set(ws, label_row, sign_col_start, "GIÁM ĐỐC", font=BOLD, align=CENTER, border=None)
    ws.merge_cells(f"{sign_col_letter_start}{label_row}:{last_col_letter}{label_row}")

    name_row = label_row + 4
    _set(ws, name_row, 2, meta.nguoi_lap_bang, font=BOLD, align=Alignment(horizontal="center"), border=None)
    if meta.giam_doc.strip():
        ws.merge_cells(f"{sign_col_letter_start}{name_row}:{last_col_letter}{name_row}")
        _set(ws, name_row, sign_col_start, meta.giam_doc.strip(), font=BOLD, align=CENTER, border=None)

    # ---- Độ rộng cột / đóng băng header --------------------------------------
    ws.freeze_panes = f"A{DATA_START_ROW}"
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions[get_column_letter(col_index["ghi_chu"])].width = 40
    ws.column_dimensions[get_column_letter(col_index["canh_bao"])].width = 40
    for c in cols:
        idx = col_index[c.key]
        letter = get_column_letter(idx)
        if letter not in ("B", get_column_letter(col_index["ghi_chu"]), get_column_letter(col_index["canh_bao"])):
            ws.column_dimensions[letter].width = 11
    ws.row_dimensions[ROW_HEADER_SUB].height = 42

    return wb
