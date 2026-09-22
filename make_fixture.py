"""Sinh 1 file M03 giả lập (đúng layout cột theo ksk_report.config) để test
pipeline end-to-end mà không cần file dữ liệu thật của Trạm Y tế."""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.utils import column_index_from_string

from ksk_report import config


def _set(ws, row: int, col_letter: str, value) -> None:
    ws.cell(row=row, column=column_index_from_string(col_letter), value=value)


def _set_group(ws, row: int, group: config.ClinicalGroup, *, chua_pht="", so_bo="", xac_dinh="", phan_loai=None, tu_choi=""):
    start = column_index_from_string(group.start_col)
    ws.cell(row=row, column=start + 0, value=chua_pht)
    ws.cell(row=row, column=start + 1, value=so_bo)
    ws.cell(row=row, column=start + 2, value=xac_dinh)
    ws.cell(row=row, column=start + 3, value=phan_loai)
    if group.has_tu_choi_kham:
        ws.cell(row=row, column=start + 4, value=tu_choi)


def build_fixture(path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = config.RAW_SHEET_NAME

    # rows 1-4: header kỹ thuật giả (không được đọc)
    for r in range(1, config.RAW_DATA_START_ROW):
        _set(ws, r, "A", f"header row {r}")

    groups_by_key = {g.key: g for g in config.CLINICAL_GROUPS}

    # --- Người 1: Nam, mọi thứ bình thường -----------------------------------
    r = config.RAW_DATA_START_ROW
    _set(ws, r, config.COL_CCCD, "079001000111")
    _set(ws, r, config.COL_HO_TEN, "Nguyễn Văn An")
    _set(ws, r, config.COL_NGAY_SINH, "01/01/1990")
    _set(ws, r, config.COL_GIOI_TINH, 1)
    _set(ws, r, config.COL_CHIEU_CAO, "170")
    _set(ws, r, config.COL_CAN_NANG, "65")
    _set(ws, r, config.COL_HA_TAM_THU, 110)
    _set(ws, r, config.COL_HA_TAM_TRUONG, 70)
    _set(ws, r, config.COL_MAT_KHONG_KINH_PHAI, "10/10")
    _set(ws, r, config.COL_MAT_KHONG_KINH_TRAI, "10/10")
    for key in groups_by_key:
        _set_group(ws, r, groups_by_key[key], phan_loai=1)
    # nam -> sản khoa/phụ khoa không tính nhưng vẫn có thể có ô trống
    _set(ws, r, config.COL_HC_SL, "4,71")
    _set(ws, r, config.COL_HGB, "14,2")
    _set(ws, r, config.COL_HCT, "42")
    _set(ws, r, config.COL_MCV, "88")
    _set(ws, r, config.COL_MCH, "28")
    _set(ws, r, config.COL_MCHC, "33")
    _set(ws, r, config.COL_RDW, "12")
    _set(ws, r, config.COL_BACH_CAU, "6,5")
    _set(ws, r, config.COL_TIEU_CAU, "250")
    _set(ws, r, config.COL_GLUCOSE, "5,10")
    _set(ws, r, config.COL_CREATININ, "80")
    _set(ws, r, config.COL_AST, "31")
    _set(ws, r, config.COL_ALT, "35")
    _set(ws, r, config.COL_TI_TRONG_NUOC_TIEU, "1,015")
    _set(ws, r, config.COL_PH_NUOC_TIEU, "6,0")
    _set(ws, r, config.COL_SIEU_AM_BUNG, "Gan, mật, tụy, lách, thận: chưa phát hiện bất thường.")
    _set(ws, r, config.COL_KET_LUAN_CHUNG, "Đủ sức khỏe lao động.")
    # Cholesterol/Triglycerid/HDL/LDL/siêu âm tuyến vú: để TRỐNG cho toàn bộ mọi
    # người trong fixture này -> kỳ vọng bị lược bỏ khỏi bảng (QUY TẮC 1)

    # --- Người 2: Nữ, ALT cao (in đậm), có chẩn đoán sơ bộ chưa xác định -----
    r += 1
    _set(ws, r, config.COL_CCCD, "079001000222")
    _set(ws, r, config.COL_HO_TEN, "Trần Thị Bình")
    _set(ws, r, config.COL_NGAY_SINH, "15/05/1985")
    _set(ws, r, config.COL_GIOI_TINH, 2)
    _set(ws, r, config.COL_CHIEU_CAO, "158")
    _set(ws, r, config.COL_CAN_NANG, "70")
    _set(ws, r, config.COL_HA_TAM_THU, 130)
    _set(ws, r, config.COL_HA_TAM_TRUONG, 85)
    _set(ws, r, config.COL_MAT_KHONG_KINH_PHAI, "9/10")
    _set(ws, r, config.COL_MAT_KHONG_KINH_TRAI, "10/10")
    for key in groups_by_key:
        _set_group(ws, r, groups_by_key[key], phan_loai=1)
    _set_group(ws, r, groups_by_key["noi_tiet"], phan_loai=2, so_bo="E11 - theo dõi ĐTĐ type 2")
    _set_group(ws, r, groups_by_key["phu_khoa"], phan_loai=2, xac_dinh="N76 - viêm âm đạo")
    _set(ws, r, config.COL_HC_SL, "4,50")
    _set(ws, r, config.COL_HGB, "13,0")
    _set(ws, r, config.COL_HCT, "40")
    _set(ws, r, config.COL_MCV, "90")
    _set(ws, r, config.COL_MCH, "29")
    _set(ws, r, config.COL_MCHC, "34")
    _set(ws, r, config.COL_RDW, "13")
    _set(ws, r, config.COL_BACH_CAU, "7,2")
    _set(ws, r, config.COL_TIEU_CAU, "230")
    _set(ws, r, config.COL_GLUCOSE, "6,80")  # cao hơn 6.20 -> in đậm
    _set(ws, r, config.COL_CREATININ, "70")
    _set(ws, r, config.COL_AST, "38")
    _set(ws, r, config.COL_ALT, "68")  # cao hơn 41 -> in đậm
    _set(ws, r, config.COL_TI_TRONG_NUOC_TIEU, "1,020")
    _set(ws, r, config.COL_PH_NUOC_TIEU, "6,5")
    _set(ws, r, config.COL_SIEU_AM_BUNG, "Gan nhiễm mỡ độ 1.")
    _set(ws, r, config.COL_KET_LUAN_CHUNG, "Cần khám chuyên khoa nội tiết.")

    # --- Người 3: Nam, Creatinin thấp (in nghiêng), dòng trống xen giữa -----
    r += 2  # cố ý để 1 dòng trống ở giữa (r+1) để test không bỏ sót người sau
    _set(ws, r, config.COL_CCCD, "079001000333")
    _set(ws, r, config.COL_HO_TEN, "Lê Văn Cường")
    _set(ws, r, config.COL_NGAY_SINH, "20/09/1978")
    _set(ws, r, config.COL_GIOI_TINH, 1)
    _set(ws, r, config.COL_CHIEU_CAO, "165")
    _set(ws, r, config.COL_CAN_NANG, "55")
    _set(ws, r, config.COL_HA_TAM_THU, 105)
    _set(ws, r, config.COL_HA_TAM_TRUONG, 65)
    _set(ws, r, config.COL_MAT_KHONG_KINH_PHAI, "10/10")
    _set(ws, r, config.COL_MAT_KHONG_KINH_TRAI, "10/10")
    for key in groups_by_key:
        _set_group(ws, r, groups_by_key[key], phan_loai=1)
    _set_group(ws, r, groups_by_key["rhm"], phan_loai=3, xac_dinh="K05 - viêm nướu")
    _set(ws, r, config.COL_HC_SL, "4,20")
    _set(ws, r, config.COL_HGB, "13,8")
    _set(ws, r, config.COL_HCT, "41")
    _set(ws, r, config.COL_MCV, "85")
    _set(ws, r, config.COL_MCH, "27")
    _set(ws, r, config.COL_MCHC, "32")
    _set(ws, r, config.COL_RDW, "12,5")
    _set(ws, r, config.COL_BACH_CAU, "5,8")
    _set(ws, r, config.COL_TIEU_CAU, "210")
    _set(ws, r, config.COL_GLUCOSE, "4,80")
    _set(ws, r, config.COL_CREATININ, "40")  # thấp hơn 46 -> in nghiêng
    _set(ws, r, config.COL_AST, "25")
    _set(ws, r, config.COL_ALT, "22")
    _set(ws, r, config.COL_TI_TRONG_NUOC_TIEU, "1,010")
    _set(ws, r, config.COL_PH_NUOC_TIEU, "9,0")  # ngoài 4.8-8.0 -> x
    _set(ws, r, config.COL_SIEU_AM_BUNG, "")
    _set(ws, r, config.COL_KET_LUAN_CHUNG, "")

    wb.save(path)
